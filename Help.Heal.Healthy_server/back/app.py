#coding=UTF-8
#This Python file uses the following encoding:utf-8
from flask import Flask,render_template,request,send_file
import pymysql
import json
import random
import datetime
import time
import openpyxl
import os 
import qrcode

#foodType {
# 1:rice
# 2:meat
# 3:vegetable
# 4:milk
# 5:nut
# 6:fruit
# }

#userName = "d0641593"
#password = "6qZb9sXTbSTN9JmC"
#databaseName = "d0641593"

userName = "iosUser"
password = "00000000"
databaseName = "iosApp"

breakfast = time.strptime("12 0 0","%H %M %S")
lunch = time.strptime("16 0 0","%H %M %S")

app = Flask(__name__)
@app.route('/register',methods = ['POST']) #註冊
def sendjson():
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            return 'Cannot post only int'
    except ValueError:
        return "Error type"
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select * from user where account = %s",jsondata['account'])
    data = cursor.fetchone()
    x = {
        "RegisterIsSuccess" : True
    }
    if(data == None) :
        cursor.execute("insert into user(account,password) value(%s,%s)",(jsondata['account'],jsondata['password']))
        db.commit()
        print("Register Success")
        db.close()
        return json.dumps(x)
    x['RegisterIsSuccess'] = False
    print("Register fail")
    db.close()
    return json.dumps(x)

@app.route('/login',methods = ['POST'])#登入
def login():
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            print("Cannot post only int")
            return 'Cannot post only int'
    except ValueError:
        print("Error type")
        return "Error type"
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select * from user where account = %s&&password = %s",(jsondata['account'],jsondata['password']))
    data = cursor.fetchone()
    x = {
        "LoginIsSuccess" : True
    }
    if(data == None) :
        x['LoginIsSuccess'] = False
        print("Login fail")
        db.close()
        return json.dumps(x)
    print("Login Success")
    db.close()
    return json.dumps(x)

@app.route('/userInfo',methods = ['POST'])#拿資料
def info() :
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            return 'Cannot post only int'
    except ValueError:
        return "Error type"
    account = jsondata["account"]
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select * from user where account = %s",account)
    data = cursor.fetchone()
    if data == None :
        print("fail to get ID from account : %s",account)
        db.close()
        return json.dumps(fail,indent =2)
    x = {
        'ID' : data[0],
        'account' : data[1],
        'password' : data[2],
        'name' : data[3],
        'sex' : data[4],
        'height' : data[5],
        'weight' : data[6],
        'birthday' : data[7]
    }
    db.close()
    return json.dumps(x,indent = 2)

@app.route('/foodInput/<account>',methods = ['POST'])#
def foodInput(account):
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            return 'Cannot post only int'
    except ValueError:
        return "Error type"
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select ID from user where account = %s",account)
    datatmp = cursor.fetchone()
    if datatmp == None : 
        db.close()
        return "fail to get ID from account : %s"%account
    userID = datatmp[0]
    foodList = jsondata["foodName"]
    foodPortionList = jsondata['foodPortion']
    print(foodList)
    timeTmp = time.strptime(jsondata["time"].replace(":"," "),"%H %M")
    if timeTmp < breakfast:
        mealTime = 1
    elif timeTmp < lunch:
        mealTime = 2
    else:
        mealTime = 3
    date = datetime.datetime.now()
    dateString = date.strftime("%Y-%m-%d")
    db_commit = 'insert into meal value(null,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)'
    for i in range(len(foodList)):
        cursor.execute("select * from food where foodName = %s",foodList[i])
        foodData = cursor.fetchone()
        foodPortion = []
        for j in range(2,len(foodData)-1):
            foodPortion.append(int(foodData[j]) * foodPortionList[i])
        print(foodPortion)
        cursor.execute(db_commit,(userID,dateString,foodList[i],foodPortion[0],foodPortion[1],foodPortion[2],foodPortion[3],foodPortion[4],foodPortion[5],mealTime))
        #userID / date / foodName / foodType / mealtime
    db.commit()
    db.close()
    return "insert success"

@app.route('/announcement')#for 公告
def announcement() :
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select * from post;")
    data = cursor.fetchall()
    dataList = list()
    for i in data :
        x = {
            "title" : i[1],
            "document" : i[2],
            "date" : i[3]
        }
        dataList.append(x)
    db.close()
    return json.dumps(dataList,indent = 2,  ensure_ascii=False).encode('utf8')

@app.route('/a_week_data',methods = ['POST'])
def forweek() :
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            return 'Cannot post only int'
    except ValueError:
        return "Error type"
    account = jsondata['account']
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute('select ID from user where account = %s',account)
    data = cursor.fetchone()
    if data == None :
        print("fail to get ID from account : %s",account)
        db.close()
        return "fail to get userID"
    userID = data[0]
    date = datetime.datetime.now()
    delay = datetime.timedelta(days = 1)
    weekdata = list()
    for i in range(7) :
        date = date - delay
        cursor.execute('select * from info where userID = %s&&date = %s',(userID,date.strftime('%Y-%m-%d')))
        data = cursor.fetchone()
        if data == None :
            x = {
                'date' : date.strftime("%m-%d"),
                'step' : -1,
                'sleep' : -1,
                'calories' : -1
            }
        else : 
            datafordate = str(data[2])
            datafordate = datafordate.replace("-","/")
            datasplit = datafordate.split("/")
            MonAndDay = str(int(datasplit[1])) + '/' + str(int(datasplit[2]))#去掉前面的零+排版
            x = {
                'date' : MonAndDay,#只要月份日期
                'calories' : data[3],
                'step' : data[4],
                'sleep' : data[5]
            }
        weekdata.append(x)
    print('success')
    db.close()
    return json.dumps(weekdata,indent =2)

@app.route('/fakedata/<account>/<date>')
def makefakedata(account,date) : 
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select ID from user where account = %s",account)
    data = cursor.fetchone()
    if data == None : 
        db.close()
        return "fail to get ID from account : %s"%account
    userID = data[0]
    calories = random.randint(1500,3500)
    step = random.randint(3000,15000)
    sleep = random.randint(3,12)
    rice = random.randint(1,5)
    meat = random.randint(2,4)
    vegetable = random.randint(2,4)
    milk = random.randint(1,3)
    nut = random.randint(1,3)
    fruit = random.randint(1,3)
    cursor.execute('insert into info value(null,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)',(userID,date,calories,step,sleep,rice,meat,vegetable,milk,nut,fruit))
    db.commit()
    db.close()
    return "insert success"

@app.route("/foodList" ,methods = ["POST"])#改post
def mealOutput():
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            return 'Cannot post only int'
    except ValueError:
        return "Error type"
    account = jsondata["account"]
    date = jsondata["date"]
    breakfastList = []
    lunchList = []
    dinnerList = []
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select ID from user where account = %s",account)
    data = cursor.fetchone()
    if data == None:
        db.close()
        return "fail to get userID from account : %s"%account
    userID = data[0]
    cursor.execute("select * from meal where userID = %s and date = %s",(userID,date))
    data = cursor.fetchall()

    if data != None:
        foodTypeLowLimit = [2,3,3,1,1,2]#定死，六大類食物每日建議攝取量
        foodTypeHighLimit = [4,8,5,2,2,4]
        foodType = ['rice','meat','vegetable','milk','nut','fruit']
        foodTypeChinese = ['穀物類','蛋豆魚肉類','蔬菜類','乳品類','堅果類','水果類']
        userFoodTypeSum = [0] * 6
        output = []
        for i in data:
            x = {
                "Date" : i[2],
                "foodName" : i[3],
                "rice" : i[4],
                "meat" : i[5],
                "vegetable" : i[6],
                "milk" : i[7],
                "nut" : i[8],
                "fruit" : i[9]
            }
            if i[10] == 1:
                breakfastList.append(x)
            elif i[10] == 2:
                lunchList.append(x)
            else:
                dinnerList.append(x)
            for j in range(6):#6大類
                userFoodTypeSum[j] = userFoodTypeSum[j] + int(i[j+4])
        print(userFoodTypeSum)
        color = []
        progressView = []
        for i in range(len(userFoodTypeSum)):
            if userFoodTypeSum[i] > foodTypeLowLimit[i]*2 :#如果大於建議攝取量的2倍
                color.append('red')
            elif userFoodTypeSum[i] > foodTypeLowLimit[i] :
                color.append('green')
            else:
                color.append('orange')
            progressView.append(float(userFoodTypeSum[i])/foodTypeHighLimit[i])
        x = {
            "progressView" : progressView,
            "color" : color,
            'image' : foodType,
            'foodType' : foodTypeChinese,
            "breakfast" : breakfastList,
            "lunch" : lunchList,
            "dinner" : dinnerList
        }
    else:
        x = {
            "progressView": [0]*6,
            "color": ['orange']*6,
            'image' : foodType,
            'foodType' : foodTypeChinese,
            "breakfast" : breakfastList,
            "lunch" : lunchList,
            "dinner" : dinnerList
        }
    db.close()
    return json.dumps(x,indent = 2,ensure_ascii=False).decode('utf8')# i don't know why is decode.
    # read this issue https://stackoverflow.com/questions/10406135/unicodedecodeerror-ascii-codec-cant-decode-byte-0xd1-in-position-2-ordinal

@app.route('/sendImage/<imageName>')
def sendImage(imageName):
    fileName = "image/" + imageName
    try :
        filePath = fileName + ".jpg"
        return send_file(filePath,mimetype = "image/jpeg")
    except :
        filePath = fileName + ".png"
        return send_file(filePath,mimetype = "image/png")

@app.route('/mealNotificationInput', methods = ["POST"])
def mealNotificationInput():
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            return 'Cannot post only int'
    except ValueError:
        return "Error type"
    account = jsondata['account']
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select ID from user where account = %s",account)
    data = cursor.fetchone()
    if data == None:
        print("fail to get userID from account = %s",account)
        db.close()
        return "fail to get userID from account = %s"%account
    userID = data[0]
    cursor.execute("select * from mealNotification where userID = %s",userID)
    data = cursor.fetchone()
    if data == None:
        cursor.execute("insert into mealNotification value(null,%s,%s,%s,%s,%s)",(userID,jsondata['title'],jsondata['body'],jsondata['hour'],jsondata["minute"]))
    else:
        cursor.execute("update mealNotification set title = %s , body = %s , hour = %s , minute = %s where userID = %s",(jsondata['title'],jsondata['body'],jsondata['hour'],jsondata["minute"],userID))
    db.commit()
    db.close()
    return "MealNotification Input Success"

@app.route('/mealNotificationOutput',methods = ["POST"])
def mealNotificationOutput():
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            return 'Cannot post only int'
    except ValueError:
        return "Error type"
    account = jsondata['account']
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select ID from user where account = %s",account)
    data = cursor.fetchone()
    if data == None:
        print("fail to get userID from account = %s",account)
        db.close()
        return "fail to get userID from account = %s"%account
    userID = data[0]
    cursor.execute("select * from mealNotification where userID = %s",userID)
    data = cursor.fetchone()
    x = {
        "title" : data[2],
        "body" : data[3],
        "hour" : data[4],
        "minute" : data[5]
    }
    db.close()
    return json.dumps(x,indent = 2)


# foodChineseName for foodRecommend
foodChineseName = ["牛排","牛肉麵","水餃","龍眼","枇杷","荔枝","東坡肉","蒸蓮藕","湯圓","炒高麗菜","炒花椰菜","牛奶","優格","優酪乳","開心果","葵花籽","杏仁"]

@app.route('/foodRecommend',methods = ['Post'])
def foodRecommend():
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            return 'Cannot post only int'
    except ValueError:
        return "Error type"
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select id from user where account = %s",jsondata['account'])
    data = cursor.fetchone()
    if data == None:
        print("fail to get userID from account = %s",account)
        db.close()
        return "fail to get userID from account = %s"%account
    userID = data[0]
    date = datetime.datetime.now().strftime('%Y-%m-%d')
    cursor.execute("select * from meal where userID = %s and date = %s",(userID,date))
    data = cursor.fetchall()
    print(data)
    foodTypeLowLimit = [2,3,3,1,1,2]#定死，六大類食物每日建議攝取量
    foodType = ['rice','meat','vegetable','milk','nut','fruit']
    userFoodTypeSum = [0] * 6
    recommendList = []
    output = []
    chineseName = []
    for i in data:
        for j in range(6):#6大類
            userFoodTypeSum[j] = userFoodTypeSum[j] + int(i[j+4])
    print(userFoodTypeSum)
    for i in range(6):#6大類
        if userFoodTypeSum[i] < foodTypeLowLimit[i]:
            cursor.execute("select foodName from food where %s >= 1"%foodType[i])
            data = cursor.fetchall()
            for j in data:
                if j[0] not in recommendList:
                    recommendList.append(j[0])
    print(recommendList)
    for i in range(3):
        if len(recommendList) == 0:
            break
        else:
            ranInt = random.randint(0,len(recommendList)-1)
            output.append(recommendList[ranInt])
            cursor.execute("select id from food where foodName = %s",output[i])
            foodID = cursor.fetchone()
            if foodID != None:
                print(foodID[0])
                chineseName.append(foodChineseName[int(foodID[0])-1])#id 比 foodChineseFood list index 大1
            recommendList.pop(ranInt)

    x = {
        'recommend' : output,
        'chineseName' : chineseName
    }
    db.close()
    return json.dumps(x,indent = 2,ensure_ascii=False).encode('utf8')

@app.route('/progressView',methods = ["POST"])
def progressView():
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            return 'Cannot post only int'
    except ValueError:
        return "Error type"
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select id from user where account = %s",jsondata['account'])
    data = cursor.fetchone()
    if data == None:
        print("fail to get userID from account = %s",account)
        db.close()
        return "fail to get userID from account = %s"%account
    userID = data[0]
    date = jsondata['date']
    print(date)
    cursor.execute("select * from meal where userID = %s and date = %s",(userID,date))
    data = cursor.fetchall()
    print(data)
    if data != None:
        foodTypeLowLimit = [2,3,3,1,1,2]#定死，六大類食物每日建議攝取量
        foodTypeHighLimit = [4,8,5,2,2,4]
        foodType = ['rice','meat','vegetable','milk','nut','fruit']
        foodTypeChinese = ['穀物類','蛋豆魚肉類','蔬菜類','乳品類','堅果類','水果類']
        userFoodTypeSum = [0] * 6
        output = []
        for i in data:
            for j in range(6):#6大類
                userFoodTypeSum[j] = userFoodTypeSum[j] + int(i[j+4])
        print(userFoodTypeSum)
        color = []
        progressView = []
        for i in range(len(userFoodTypeSum)):
            if userFoodTypeSum[i] > foodTypeLowLimit[i]*2 :#如果大於建議攝取量的2倍
                color.append('red')
            elif userFoodTypeSum[i] > foodTypeLowLimit[i] :
                color.append('green')
            else:
                color.append('orange')
            progressView.append(float(userFoodTypeSum[i])/foodTypeHighLimit[i])
        x = {
            "progressView" : progressView,
            "color" : color,
            'image' : foodType,
            'foodType' : foodTypeChinese
        }
    else:
        x = {
            "progressView": [0]*6,
            "color": ['orange']*6,
            'image' : foodType,
            'foodType' : foodTypeChinese
        }
    db.close()
    return json.dumps(x,indent = 2)
keys = []
@app.route('/makeqrcode',methods=['POST'])#qrcode現在沒用
def makeqrcode():
    data = request.get_data()
    print(data)
    try:
        jsondata = json.loads(data)
        if(type(jsondata) == int):
            return 'Cannot post only int'
    except ValueError:
        return "Error type"
    account = jsondata['account']
    qr = qrcode.QRCode()
    key = random.randint(10000,99999)
    key = str(key)
    while key in keys:
        key = random.randint(10000,99999)
        key = str(key)
    keys.append(key)
    keys.append(account)
    print(keys)
    url = "http://sunflower.hj-deal.science/userInfoOutput/" + key
    qr.add_data(url)
    qr.make(fit = 1)
    img = qr.make_image()
    fileName = account + ".jpg"
    img.save(fileName,"JPEG")
    return send_file(fileName,mimetype = "image/jpeg")

@app.route('/userInfoOutput/<key>')
def userInfoOutput(key):
    if key in keys:
        keyIndex = keys.index(key)
        print(keyIndex)
        account = keys[(keyIndex+1)]
        keys.pop(keyIndex+1)
        keys.pop(keyIndex)#跟上面不可以相反
    else:
        return 'fail to get excel from this key'
    db = pymysql.connect('localhost',userName,password,databaseName)
    cursor = db.cursor()
    cursor.execute("select id from user where account = %s",account)
    data = cursor.fetchone()
    if data == None:
        print("fail to get userID from account = %s",account)
        db.close()
        return "fail to get userID from account = %s"%account
    userID = data[0]
    filePath = "excels/userOutputForAccount" + account + '.xlsx'
    date = datetime.datetime.now()
    delay = datetime.timedelta(days = 1)
    delay7day = datetime.timedelta(days = 7)
    date = date - delay7day
    print(date)
    if os.path.exists(filePath):
        os.remove(filePath)
    wb = openpyxl.Workbook()
    ws = wb.active
    foodTypeforxlsx = ['日期','卡路里','步數','睡眠','穀物類','蛋豆魚肉類','蔬菜類','奶類','堅果類','水果類']
    ws.append(foodTypeforxlsx)
    for i in range(7): #7天
        cursor.execute("select * from meal where userID = %s and date = %s",(userID,date.strftime("%Y-%m-%d")))
        data = cursor.fetchall()
        userFoodTypeSum = [0] * 6
        if data != None:
            for x in data:
                for j in range(6):#6大類
                    userFoodTypeSum[j] = userFoodTypeSum[j] + int(x[j+4])
        cursor.execute("select * from info where userID = %s and date = %s",(userID,date.strftime("%Y-%m-%d")))
        userInfo = cursor.fetchone()
        userInfoList = []
        if userInfo != None: 
            for y in range(3,6):
                    userInfoList.append(userInfo[y])
        else:
            userInfoList = [0]*3
        strDate = date.strftime('%Y-%m-%d')
        userSumForXlsx = [strDate]
        userSumForXlsx.extend(userInfoList)
        userSumForXlsx.extend(userFoodTypeSum)
        ws.append(userSumForXlsx)
        date = date + delay
    chart = openpyxl.chart.LineChart()
    values = openpyxl.chart.Reference(ws, min_col = 5, min_row = 1, max_col = 10, max_row = 8) 
    chart.add_data(values,titles_from_data = True)
    dates = openpyxl.chart.Reference(ws, min_col=1, min_row=2, max_row=8)
    chart.set_categories(dates)
    chart.title = "Line Diagram"
    chart.x_axis.title = "Date"
    chart.y_axis.title = "Quantity"
    ws.add_chart(chart, 'A25') 
    infoChart = openpyxl.chart.LineChart()
    infoValue = openpyxl.chart.Reference(ws,min_col = 2,min_row = 1, max_col = 4,max_row = 8)
    infoChart.add_data(infoValue,titles_from_data = True)
    infoChart.set_categories(dates)
    infoChart.title = "Line Diagram"
    infoChart.x_axis.title = "Date"
    infoChart.y_axis.title = "Value"
    ws.add_chart(infoChart,'A10')
    wb.save(filePath)
    fileName = "UserInfo.xlsx"
    return send_file(filePath,as_attachment = True,attachment_filename = fileName,mimetype = "text/xlsx")


    
    


if __name__ == '__main__':
    app.run()
# flask run --host=0.0.0.0 --port=3000
